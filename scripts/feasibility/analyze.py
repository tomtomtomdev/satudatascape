import json, sys, io, csv, collections, zipfile, warnings
warnings.filterwarnings("ignore")
R = json.load(open(sys.argv[1]))
def sniff(b):
    h = b[:8]; s = b[:2048].lstrip().lower()
    if h.startswith(b"%PDF"): return "PDF"
    if h.startswith(b"PK"):
        try:
            n = zipfile.ZipFile(io.BytesIO(b)).namelist()
            return "XLSX" if any(x.startswith("xl/") for x in n) else "DOCX" if any(x.startswith("word/") for x in n) else "ZIP"
        except Exception: return "ZIP?"
    if h.startswith(b"\xd0\xcf\x11\xe0"): return "XLS"
    if s.startswith((b"<!doctype html", b"<html")) or b"<html" in s[:500]: return "HTML"
    if s.startswith((b"<?xml", b"<wms", b"<wfs", b"<ogc", b"<servicee")): return "XML"
    if s.startswith((b"{", b"[")): return "JSON"
    if h[:3] in (b"\xff\xd8\xff",) or h.startswith(b"\x89PNG"): return "IMG"
    if not b.strip(): return "EMPTY"
    return "TEXT"
def xlsx(b):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=False, data_only=True)
    ws = wb.worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True)]
    nonempty = [r for r in rows if any(c not in (None, "") for c in r)]
    hdr_idx = next((i for i, r in enumerate(rows) if sum(c not in (None, "") for c in r) >= max(2, 0.6 * max(1, max((sum(c not in (None,"") for c in x) for x in rows), default=1)))), None)
    merged = len(ws.merged_cells.ranges); top_merged = sum(1 for m in ws.merged_cells.ranges if m.min_row <= (hdr_idx or 0) + 2)
    nums = sum(isinstance(c, (int, float)) for r in nonempty for c in r); cells = sum(c not in (None, "") for r in nonempty for c in r)
    return dict(sheets=len(wb.worksheets), rows=len(nonempty), cols=ws.max_column, header_row=hdr_idx, merged=merged, top_merged=top_merged, numeric_frac=round(nums / max(cells, 1), 2),
                tidy=(hdr_idx == 0 and merged == 0 and len(wb.worksheets) == 1))
def xls(b):
    import xlrd
    wb = xlrd.open_workbook(file_contents=b, formatting_info=False); sh = wb.sheet_by_index(0)
    return dict(sheets=wb.nsheets, rows=sh.nrows, cols=sh.ncols, merged=len(getattr(sh, "merged_cells", [])))
def csvp(b):
    for enc in ("utf-8-sig", "cp1252"):
        try: t = b.decode(enc); break
        except UnicodeDecodeError: continue
    try: d = csv.Sniffer().sniff(t[:5000], delimiters=",;\t|").delimiter
    except Exception: d = ","
    rows = list(csv.reader(io.StringIO(t), delimiter=d))
    widths = collections.Counter(len(r) for r in rows if r)
    mode, cnt = widths.most_common(1)[0] if widths else (0, 0)
    return dict(enc=enc, delim=d, rows=len(rows), cols=mode, consistent=round(cnt / max(len([r for r in rows if r]), 1), 2), header=rows[0][:6] if rows else None)
def pdf(b):
    import pdfplumber
    with pdfplumber.open(io.BytesIO(b)) as p:
        n = len(p.pages); pg = p.pages[: min(3, n)]
        chars = sum(len(x.extract_text() or "") for x in pg); tables = sum(len(x.find_tables()) for x in pg)
    return dict(pages=n, chars_per_page=chars // max(len(pg), 1), tables_first3=tables, scanned=chars < 50 * len(pg))
def jsn(b):
    d = json.loads(b.decode("utf-8-sig"))
    def shape(x, depth=0):
        if isinstance(x, list): return f"list[{len(x)}]" + ("<" + shape(x[0], depth+1) + ">" if x and depth < 2 else "")
        if isinstance(x, dict): return "{" + ",".join(list(x)[:6]) + "}"
        return type(x).__name__
    recs = d if isinstance(d, list) else next((v for v in d.values() if isinstance(v, list)), None) if isinstance(d, dict) else None
    return dict(shape=shape(d)[:150], records=len(recs) if recs is not None else None, flat=bool(recs) and isinstance(recs[0], dict) and all(not isinstance(v, (dict, list)) for v in recs[0].values()))
out = []
for r in R:
    a = dict(fmt=r["fmt"], fmt_raw=r.get("fmt_raw"), host=r.get("host"), org=r.get("org"), status=r.get("status"), err=r.get("err"), bytes=r.get("bytes"), trunc=r.get("trunc"), tls=r.get("tls"), ctype=r.get("ctype"), url=r.get("url"), secs=r.get("secs"))
    if r.get("status") == 200 and r.get("path"):
        b = open(r["path"], "rb").read(); a["actual"] = sniff(b)
        fn = {"XLSX": xlsx, "XLS": xls, "PDF": pdf, "JSON": jsn}.get(a["actual"])
        if a["actual"] == "TEXT": fn = csvp
        if fn and not r.get("trunc"):
            try: a["parse"] = fn(b); a["parse_ok"] = True
            except Exception as e: a["parse_ok"] = False; a["parse_err"] = type(e).__name__ + ":" + str(e)[:100]
        if a["actual"] == "HTML": a["html_title"] = (b.split(b"<title>")[1].split(b"</title>")[0][:80].decode("utf-8", "replace") if b"<title>" in b else None)
        if a["actual"] == "XML": a["xml_head"] = b[:200].decode("utf-8", "replace")
    out.append(a)
json.dump(out, open(sys.argv[2], "w"), indent=1, default=str)
print("done", len(out))
